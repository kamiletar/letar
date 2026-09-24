# Генератор XLSX-пакета ревьюеру по части B аудита банка (этап 5.10).
#
# Источник — машинный отчёт docs/question-bank-audit.json (часть A) и сам дамп
# prisma/questions-dump.json. В отличие от generate-question-review-xlsx.py, здесь не новые
# вопросы из батча, а УЖЕ работающие в тесте, отобранные аудитом как подозрительные:
#   1. «Расхождения» — 11 известных KNOWN_DISCREPANCIES: справочник max-баллов заявляет
#      шкалы, которых нет в вариантах.
#   2. «Дубли» — почти-дословные пары (≥ 0.65) и сильные сюжетные (≥ 0.8).
#   3. «Вариант-вездеход» — вариант, скорящий 5+ шкал.
# Редких шкал (SAD, ASD, MAS, ALX) здесь нет: все их вопросы пришли батчем 5.1 и уже лежат
# в question-review-5.1.xlsx (проверка ниже падает, если это перестанет быть правдой).
#
# Запуск (из корня apps/archetest): python scripts/generate-audit-review-xlsx.py
# Требует: pip install openpyxl

import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

APP_DIR = Path(__file__).resolve().parent.parent
OUT_FILE = APP_DIR / "docs" / "question-review" / "question-review-audit-b.xlsx"

# Общие стили — из генератора основных пакетов (имя файла с дефисами → через importlib)
_spec = importlib.util.spec_from_file_location(
    "review_xlsx", Path(__file__).resolve().parent / "generate-question-review-xlsx.py"
)
review_xlsx = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(review_xlsx)
FONT, HEADER_FILL, BORDER = review_xlsx.FONT, review_xlsx.HEADER_FILL, review_xlsx.BORDER
fmt_scoring, build_intro_sheet = review_xlsx.fmt_scoring, review_xlsx.build_intro_sheet

RARE_SCALES = ["SAD", "ASD", "MAS", "ALX"]
NEAR_DUP_MIN = 0.65
STORY_DUP_MIN = 0.8

INTRO = [
    ("Ревью вопросов Архетеста — часть B аудита банка: вопросы, которые УЖЕ работают в тесте", True),
    ("", False),
    ("Машинный аудит прошёл по всем 2126 вопросам и отобрал подозрительные. Ручная вычитка всего", False),
    ("банка — это 35–50 часов, поэтому здесь только то, куда стоит направить внимание. Срока нет.", False),
    ("", False),
    ("Листы — в порядке важности:", True),
    ("1. «Расхождения» — 11 вопросов, где справочник максимальных баллов заявляет шкалы, которых нет", False),
    ("   в вариантах ответа. Нужно решение: добавить балл в вариант или убрать шкалу из справочника.", False),
    ("2. «Дубли» — пары вопросов с одним сюжетом. Решение: оставить оба / убрать один / переписать.", False),
    ("3. «Вездеход» — вариант, который начисляет баллы пяти шкалам сразу: не размыт ли он?", False),
    ("", False),
    ("Редкие шкалы (садизм, систематизация, мазохизм, алекситимия — меньше 30 вопросов каждая):", True),
    ("все их вопросы уже в пакете 5.1. Если будете смотреть 5.1 — начните с этих четырёх шкал:", False),
    ("там каждый вопрос весит в разы больше обычного (по садизму их 21 на всю шкалу).", False),
    ("", False),
    ("⚠️ Важно: правка БАЛЛОВ работающего вопроса делает старые результаты несопоставимыми с новыми.", True),
    ("Поэтому «Править формулировку» и «Править баллы» — разные вердикты. Первый применяем сразу,", False),
    ("второй копим и применяем пачкой одним осознанным обновлением версии банка.", False),
    ("", False),
    ("Критерии — те же, что в INSTRUCTIONS.md (конструкт, язык, социальная желательность, этика).", False),
]

VERDICT_QUESTION = '"ОК,Править формулировку,Править баллы,Удалить"'
VERDICT_PAIR = '"Оставить оба,Убрать A,Убрать B,Переписать один"'


def style_sheet(ws, headers, widths, verdict_col, verdict_list):
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    if ws.max_row >= 2:
        dv = DataValidation(type="list", formula1=verdict_list, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{verdict_col}2:{verdict_col}{ws.max_row}")
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last}{ws.max_row}"
    ws.freeze_panes = "C2"


def option_cells(q):
    cells = [f"{o['text']}\n[{fmt_scoring(o['scoring'])}]" for o in q["options"]]
    return (cells + [""] * 4)[:4]


def main():
    audit = json.loads((APP_DIR / "docs" / "question-bank-audit.json").read_text(encoding="utf-8"))
    dump = json.loads((APP_DIR / "prisma" / "questions-dump.json").read_text(encoding="utf-8"))
    by_qnum = {}
    for q in dump:
        by_qnum[q["sortOrder"] + 1] = {**q, "options": json.loads(q["options"])}

    batch_51 = set()
    for path in (APP_DIR / "prisma" / "question-batches" / "5.1").glob("*.json"):
        batch_51 |= {q["scenario"] for q in json.loads(path.read_text(encoding="utf-8"))}

    wb = Workbook()
    intro = wb.active
    intro.title = "Инструкция"
    build_intro_sheet(intro, INTRO)

    q_headers = ["№", "Шкалы", "Сценарий", "Вариант А", "Вариант Б", "Вариант В", "Вариант Г", "Вердикт", "Комментарий / правка"]
    q_widths = [6, 14, 45, 34, 34, 34, 34, 16, 45]

    # Инвариант, на котором держится отсутствие листа «Редкие шкалы»
    for qnum, q in by_qnum.items():
        if q["scenario"] not in batch_51 and any(c in o["scoring"] for o in q["options"] for c in RARE_SCALES):
            raise SystemExit(f"№{qnum}: вопрос редкой шкалы вне батча 5.1 — верни лист «Редкие шкалы»")

    disc = wb.create_sheet("Расхождения")
    d_headers = ["№", "Справочник заявляет (шкала: в вариантах → в справочнике)", "Сценарий", "Вариант А", "Вариант Б", "Вариант В", "Вариант Г", "Вердикт", "Комментарий / правка"]
    disc.append(d_headers)
    for item in audit["discrepancies"]["known"]:
        q = by_qnum[item["qnum"]]
        diff = ", ".join(f"{code}: {a} → {b}" for code, (a, b) in item["diff"].items())
        disc.append([item["qnum"], diff, q["scenario"], *option_cells(q), "", ""])
    style_sheet(disc, d_headers, [6, 26, 45, 34, 34, 34, 34, 16, 45], "H", VERDICT_QUESTION)

    dups = wb.create_sheet("Дубли")
    p_headers = ["№ A", "№ B", "Тип", "Сходство", "Сценарий A", "Сценарий B", "Вердикт", "Комментарий"]
    dups.append(p_headers)
    pairs = [("почти дословно", x) for x in audit["nearDuplicates"] if x["similarity"] >= NEAR_DUP_MIN]
    pairs += [("сюжет", x) for x in audit["storyDuplicates"] if x["similarity"] >= STORY_DUP_MIN]
    seen = set()
    for kind, x in pairs:
        key = (min(x["a"], x["b"]), max(x["a"], x["b"]))
        if key in seen:
            continue
        seen.add(key)
        dups.append([x["a"], x["b"], kind, round(x["similarity"], 2), x["scenarioA"], x["scenarioB"], "", ""])
    style_sheet(dups, p_headers, [7, 7, 14, 9, 50, 50, 16, 40], "G", VERDICT_PAIR)

    fat = wb.create_sheet("Вездеход")
    fat.append(q_headers)
    for item in audit["asymmetry"]["fatOptions"]:
        q = by_qnum[item["qnum"]]
        fat.append([item["qnum"], f"вариант {'АБВГ'[item['optionIndex']]}: {item['scales']} шкал", q["scenario"], *option_cells(q), "", ""])
    style_sheet(fat, q_headers, q_widths, "H", VERDICT_QUESTION)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(
        f"OK: {OUT_FILE} — расхождения {disc.max_row - 1}, "
        f"пары дублей {dups.max_row - 1}, вездеход {fat.max_row - 1}"
    )


if __name__ == "__main__":
    main()
